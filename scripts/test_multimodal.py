import os
import sys
import docx
import asyncio
import logging
from httpx import AsyncClient, ASGITransport
from sqlalchemy import select

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from api.main import app
from db.session import init_db, async_session_factory
from db.models import Document, Chunk
from ingestion.pipeline import IngestionPipeline
from api.rag.bm25_index import bm25_index

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(name)s: %(message)s")
logger = logging.getLogger("test_multimodal")

def create_sample_multimodal_docx(target_path: str):
    os.makedirs(os.path.dirname(target_path), exist_ok=True)
    doc = docx.Document()
    doc.add_heading("Chapter 4: AVL Trees and Comparative Complexity", level=1)

    p1 = doc.add_paragraph(
        "In strictly balanced trees, the balance factor for any node is defined as "
        "$BF(v) = h(v.left) - h(v.right)$, where $h$ denotes subtree height. "
        "Every valid AVL node strictly preserves the condition $-1 \\le BF(v) \\le 1$."
    )

    # Add Academic Data Table
    doc.add_heading("Algorithm Complexity Comparison Table", level=2)
    table = doc.add_table(rows=4, cols=4)
    headers = ["Data Structure", "Average Search", "Worst-Case Search", "Space Complexity"]
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h

    data_rows = [
        ["Unsorted Array", "O(n)", "O(n)", "O(n)"],
        ["Sorted Array", "O(log n)", "O(log n)", "O(n)"],
        ["AVL Balanced Tree", "O(log n)", "O(log n)", "O(n)"]
    ]
    for row_idx, row_values in enumerate(data_rows):
        for col_idx, val in enumerate(row_values):
            table.rows[row_idx + 1].cells[col_idx].text = val

    # Add Diagram / Figure Caption
    doc.add_paragraph("Figure 1: AVL Tree Left-Right Double Rotation State Machine Diagram.")
    doc.save(target_path)
    logger.info(f"Created sample multi-modal document at: {target_path}")

async def run_multimodal_test():
    await init_db()

    target_docx = os.path.abspath("data/sample_courses/CS401/multimodal_lecture.docx")
    create_sample_multimodal_docx(target_docx)

    # 1. Ingest document
    logger.info("--- 1. Ingesting Multi-Modal Document ---")
    async with async_session_factory() as session:
        res = await IngestionPipeline.process_file(target_docx, session)
        assert res["status"] in ("success", "indexed", "unchanged"), f"Ingestion failed: {res}"
        logger.info(f"Ingestion status: {res['status']}")

    # 2. Inspect Extracted Chunks
    logger.info("\n--- 2. Inspecting Multi-Modal Chunk Extractions ---")
    async with async_session_factory() as session:
        stmt = (
            select(Chunk)
            .join(Document)
            .where(Document.title == "multimodal_lecture.docx")
        )
        chunks = (await session.execute(stmt)).scalars().all()
        assert len(chunks) > 0, "No chunks found for multimodal_lecture.docx"

        full_extracted_text = "\n".join(c.text for c in chunks)

        # Check Table Serialization
        assert "| Data Structure | Average Search |" in full_extracted_text, "Markdown table header not found!"
        assert "| AVL Balanced Tree |" in full_extracted_text, "Table rows not found!"
        logger.info("✓ Structured Markdown Table serialized with row/column alignment.")

        # Check Equation Preservation
        assert "$BF(v) = h(v.left) - h(v.right)$" in full_extracted_text, "LaTeX formula was stripped or corrupted!"
        logger.info("✓ LaTeX mathematical formula preserved in chunk text.")

        # Check Figure Caption
        assert "**[Figure/Diagram]** Figure 1: AVL Tree Left-Right" in full_extracted_text, "Figure caption tag missing!"
        logger.info("✓ Multi-modal diagram caption tagged and indexed.")

    # 3. Rebuild BM25 Index to include new chunks
    bm25_index.corpus = []
    bm25_index.ensure_loaded()
    logger.info(f"Rebuilt BM25 index with {len(bm25_index.corpus)} total chunks.")

    # 4. End-to-End Query Verification via API
    logger.info("\n--- 3. Testing Retrieval & Grounded Answering ---")
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://testserver") as client:
        # Query asking for table data
        table_q = "What is the worst-case search complexity of an AVL Balanced Tree in the comparison table?"
        res = await client.post("/api/v1/ask", json={
            "question": table_q,
            "department": "ComputerScience",
            "course": "CS401"
        })
        assert res.status_code == 200
        data = res.json()
        assert len(data["citations"]) > 0
        assert any("multimodal_lecture.docx" in c["title"] for c in data["citations"])
        logger.info(f"✓ Table Query Answer: '{data['answer'][:80]}...' (Citations: {len(data['citations'])})")

        # Query asking for mathematical formula
        math_q = "What is the formula for the balance factor BF(v) of an AVL tree node?"
        res_math = await client.post("/api/v1/ask", json={
            "question": math_q,
            "department": "ComputerScience",
            "course": "CS401"
        })
        assert res_math.status_code == 200
        math_data = res_math.json()
        assert len(math_data["citations"]) > 0
        assert any("multimodal_lecture.docx" in c["title"] for c in math_data["citations"])
        logger.info(f"✓ Math Formula Answer: '{math_data['answer'][:80]}...' (Citations: {len(math_data['citations'])})")

    logger.info("\n=======================================================")
    logger.info("ALL PHASE 16 MULTI-MODAL INGESTION TESTS PASSED!")
    logger.info("=======================================================")

if __name__ == "__main__":
    asyncio.run(run_multimodal_test())
