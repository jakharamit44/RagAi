import docx

doc = docx.Document("Enterprise_RAG_University_Plan_September_2026_Updated.docx")
print("Paragraphs containing principles:")
for i, p in enumerate(doc.paragraphs):
    if "principle" in p.text.lower():
        print(f"  Para {i} (style={getattr(p.style, 'name', '')}): {p.text}")

print("\nTables containing principles:")
for t_idx, t in enumerate(doc.tables):
    for r_idx, r in enumerate(t.rows):
        for c_idx, c in enumerate(r.cells):
            if "principle" in c.text.lower():
                print(f"  Table {t_idx} Row {r_idx} Col {c_idx}: {c.text[:100]}")
