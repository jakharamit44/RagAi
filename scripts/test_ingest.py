import os
import sys
import asyncio
import logging
import docx

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from db.session import init_db, async_session_factory
from db.models import ManifestEntry, Document, Chunk
from ingestion.folder_watcher import FolderWatcher
from sqlalchemy import select

logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")
logger = logging.getLogger("test_ingest")

def create_sample_docx(file_path: str):
    doc = docx.Document()
    doc.add_heading("CS401: Advanced Data Structures & Algorithms", level=1)
    doc.add_paragraph("Welcome to CS401. This course covers balanced search trees, graph algorithms, and dynamic programming.")

    doc.add_heading("Chapter 1: AVL Trees", level=2)
    doc.add_paragraph("An AVL tree is a self-balancing binary search tree. In an AVL tree, the heights of the two child subtrees of any node differ by at most one.")
    doc.add_paragraph("If at any time the heights differ by more than one, rebalancing is done to restore this property through tree rotations.")

    doc.add_heading("Chapter 2: Graph Traversal", level=2)
    doc.add_paragraph("Breadth-first search (BFS) traverses a tree or graph level by level. It uses a queue data structure with FIFO ordering.")
    doc.add_paragraph("Depth-first search (DFS) explores as far as possible along each branch before backtracking. It uses a stack or recursion.")

    os.makedirs(os.path.dirname(file_path), exist_ok=True)
    doc.save(file_path)

def create_sample_txt(file_path: str):
    os.makedirs(os.path.dirname(file_path), exist_ok=True)
    with open(file_path, "w", encoding="utf-8") as f:
        f.write("# CS401 Syllabus\n\nInstructor: Dr. Alan Turing\nOffice Hours: Mon/Wed 2-4 PM\nLocation: Turing Hall Room 302\n\nPrerequisites: CS201 Data Structures and CS202 Discrete Mathematics.")

async def run_verification():
    await init_db()

    # 1. Create sample course hierarchy
    test_root = os.path.abspath("data/sample_courses")
    course_dir = os.path.join(test_root, "ComputerScience", "Semester4", "CS401")
    docx_file = os.path.join(course_dir, "lecture1_notes.docx")
    txt_file = os.path.join(course_dir, "course_syllabus.txt")

    create_sample_docx(docx_file)
    create_sample_txt(txt_file)
    logger.info(f"Created sample course documents in: {course_dir}")

    watcher = FolderWatcher()

    # 2. Run initial reconciliation scan
    logger.info("--- PASS 1: Initial Ingestion ---")
    results1 = await watcher.reconcile_folder(test_root)

    for r in results1:
        logger.info(f"File processed: {r.get('title', r.get('path'))} -> status={r.get('status')}, chunks={r.get('chunks_count', 0)}")

    # Verify DB records
    async with async_session_factory() as session:
        manifest_res = await session.execute(select(ManifestEntry))
        manifests = manifest_res.scalars().all()
        doc_res = await session.execute(select(Document))
        docs = doc_res.scalars().all()
        chunk_res = await session.execute(select(Chunk))
        chunks = chunk_res.scalars().all()

        logger.info(f"DB Manifest Entries: {len(manifests)}")
        for m in manifests:
            logger.info(f"  Manifest: {os.path.basename(m.path)} | hash={m.content_hash[:10]}... | status={m.status}")

        logger.info(f"DB Documents: {len(docs)}")
        for d in docs:
            logger.info(f"  Doc: {d.title} | dept={d.department} | sem={d.semester} | course={d.course} | type={d.doc_type}")

        logger.info(f"DB Chunks: {len(chunks)}")
        for c in chunks[:3]:
            logger.info(f"  Chunk: sec='{c.section}' | page={c.page_number} | text='{c.text[:60]}...'")

        cs401_docs = [d for d in docs if d.course == "CS401"]
        assert len(cs401_docs) >= 2, "Expected at least 2 ingested CS401 documents"
        assert len(chunks) >= 3, "Expected at least 3 chunks"
        assert all(d.department == "ComputerScience" for d in cs401_docs), "Path tagging failed for department"

    # 3. Run second pass to verify idempotency (skip unchanged files)
    logger.info("--- PASS 2: Idempotent Rescan (Expect skips) ---")
    results2 = await watcher.reconcile_folder(test_root)
    for r in results2:
        logger.info(f"Rescan result: {r.get('path', '')} -> status={r.get('status')}, reason={r.get('reason')}")
        assert r.get("status") == "skipped", "Unchanged file was not skipped!"

    logger.info("ALL VERIFICATIONS PASSED SUCCESSFULLY! Phase 1 Ingestion works cleanly.")

if __name__ == "__main__":
    asyncio.run(run_verification())
